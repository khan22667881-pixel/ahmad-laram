"""ATS Resume Checker: upload a resume, get an estimated ATS score and fixes.

UI: Streamlit    AI: Google Gemini Flash (via the `google-genai` SDK)

Note: the score is an AI-based *estimate* of how well a resume would parse and
match in typical applicant tracking systems. It is not an official score from
any real ATS.
"""

import io
import json
import os
import re

import streamlit as st
# docx, pypdf and google-genai are imported lazily (inside the functions that use them)
# so a missing dependency shows a clear message in the app instead of a crash.

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
DEFAULT_MODEL = "gemini-2.5-flash"  # override with GEMINI_MODEL (secret/env) or the sidebar
MAX_FILE_MB = 5
MAX_CHARS = 30_000  # resume text sent to the model (keeps cost and latency low)
MIN_CHARS = 200  # below this we assume the file has no readable text


MISSING_DEPS_MSG = (
    "A required package is not installed ({err}). Make sure `requirements.txt` is in the "
    "ROOT of your GitHub repo (same folder as app.py), is named exactly `requirements.txt`, "
    "is saved as UTF-8, contains streamlit, google-genai, pypdf and python-docx, then reboot "
    "the app from 'Manage app'."
)


# --------------------------------------------------------------------------- #
# Helpers: configuration
# --------------------------------------------------------------------------- #
def get_config(name: str, default: str = "") -> str:
    """Read a setting from Streamlit secrets first, then environment variables."""
    try:
        value = st.secrets[name]
        if value:
            return str(value)
    except Exception:  # no secrets file, or key missing
        pass
    return os.getenv(name, default)


# --------------------------------------------------------------------------- #
# Helpers: text extraction
# --------------------------------------------------------------------------- #
def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("This PDF is password-protected.") from exc
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_docx(data: bytes) -> str:
    from docx import Document  # provided by the `python-docx` package

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many resume templates put content in tables
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Return cleaned plain text from a PDF, DOCX or TXT resume."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        text = extract_pdf(data)
    elif ext == "docx":
        text = extract_docx(data)
    elif ext == "txt":
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- #
# Helpers: prompting and parsing
# --------------------------------------------------------------------------- #
SYSTEM_INSTRUCTION = """You are an expert technical recruiter and Applicant Tracking System (ATS) \
specialist. You evaluate resumes strictly and consistently. Respond with JSON only."""

JSON_SHAPE = """{
  "overall_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "category_scores": [
    {"category": "Formatting & Parseability", "score": <0-100>, "feedback": "<1-2 sentences>"},
    {"category": "Keywords & Skills", "score": <0-100>, "feedback": "<1-2 sentences>"},
    {"category": "Impact & Achievements", "score": <0-100>, "feedback": "<1-2 sentences>"},
    {"category": "Structure & Sections", "score": <0-100>, "feedback": "<1-2 sentences>"},
    {"category": "Readability & Language", "score": <0-100>, "feedback": "<1-2 sentences>"}
  ],
  "strengths": ["<short bullet>", "..."],
  "improvements": [
    {"priority": "High" | "Medium" | "Low", "issue": "<what is wrong>", "fix": "<specific action>"}
  ],
  "missing_keywords": ["<keyword>", "..."],
  "rewrite_examples": [
    {"original": "<weak line copied from the resume>", "improved": "<stronger version>"}
  ]
}"""


def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = job_description.strip()
    if jd:
        jd_block = (
            "A target job description is provided. Add a sixth category_scores entry named "
            '"Job Match" scoring how well the resume matches it, and make "missing_keywords" '
            "the important keywords from the job description that the resume lacks.\n\n"
            f"JOB DESCRIPTION:\n\"\"\"\n{jd[:10_000]}\n\"\"\"\n"
        )
    else:
        jd_block = (
            "No job description is provided. Judge against general ATS best practices and make "
            '"missing_keywords" the commonly expected skills/keywords for the role the resume '
            "appears to target.\n"
        )

    return f"""Evaluate the resume below for ATS compatibility and quality.

Scoring rules:
- Be strict and realistic. Average resumes score 55-70; reserve 85+ for excellent ones.
- Penalize: missing contact info, missing sections (summary, experience, education, skills),
  vague duties with no measurable results, keyword gaps, inconsistent dates, typos,
  text that looks garbled by tables, columns or graphics, and excessive length.
- Reward: clear standard headings, quantified achievements, strong action verbs,
  relevant keywords, consistent formatting.
- Give 4-8 improvements, ordered High to Low priority, each specific and actionable.
- Give up to 10 missing_keywords and 2-4 rewrite_examples. Only use lines that actually
  appear in the resume for "original". Never invent experience the candidate does not have.

{jd_block}
Return ONLY valid JSON in exactly this shape:
{JSON_SHAPE}

RESUME:
\"\"\"
{resume_text[:MAX_CHARS]}
\"\"\"
"""


def _clamp(value, low=0, high=100) -> int:
    try:
        return max(low, min(high, int(round(float(value)))))
    except (TypeError, ValueError):
        return low


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def parse_analysis(raw: str) -> dict:
    """Turn the model's reply into a clean, validated dict (raises ValueError if unusable)."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start == -1 or end <= start:
            raise ValueError("The AI did not return valid JSON.")
        try:
            data = json.loads(text[start : end + 1])
        except json.JSONDecodeError as exc:
            raise ValueError("The AI returned malformed JSON.") from exc
    if not isinstance(data, dict):
        raise ValueError("The AI returned an unexpected format.")

    categories = []
    for item in data.get("category_scores") or []:
        if isinstance(item, dict) and item.get("category"):
            categories.append(
                {
                    "category": str(item["category"]),
                    "score": _clamp(item.get("score")),
                    "feedback": str(item.get("feedback", "")).strip(),
                }
            )

    if "overall_score" in data:
        overall = _clamp(data["overall_score"])
    elif categories:
        overall = round(sum(c["score"] for c in categories) / len(categories))
    else:
        raise ValueError("The AI response had no score.")

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict) and (item.get("issue") or item.get("fix")):
            priority = str(item.get("priority", "Medium")).strip().capitalize()
            improvements.append(
                {
                    "priority": priority if priority in ("High", "Medium", "Low") else "Medium",
                    "issue": str(item.get("issue", "")).strip(),
                    "fix": str(item.get("fix", "")).strip(),
                }
            )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda i: order[i["priority"]])

    rewrites = []
    for item in data.get("rewrite_examples") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append(
                {"original": str(item["original"]).strip(), "improved": str(item["improved"]).strip()}
            )

    return {
        "overall_score": overall,
        "summary": str(data.get("summary", "")).strip(),
        "category_scores": categories,
        "strengths": _str_list(data.get("strengths")),
        "improvements": improvements,
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "rewrite_examples": rewrites,
    }


def analyze_resume(client, model: str, resume_text: str, job_description: str = "") -> dict:
    """Call Gemini and return the parsed analysis dict."""
    from google.genai import types

    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_INSTRUCTION,
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return parse_analysis(response.text)


# --------------------------------------------------------------------------- #
# UI
# --------------------------------------------------------------------------- #
def score_label(score: int) -> str:
    if score >= 85:
        return "Excellent"
    if score >= 70:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    score = result["overall_score"]
    left, right = st.columns([1, 3])
    left.metric("ATS Score", f"{score}/100", score_label(score), delta_color="off")
    right.progress(score / 100)
    if result["summary"]:
        right.write(result["summary"])

    tab_scores, tab_fix, tab_keys, tab_rewrite = st.tabs(
        ["Score breakdown", "Improvements", "Keywords", "Rewrite examples"]
    )

    with tab_scores:
        for cat in result["category_scores"]:
            st.markdown(f"**{cat['category']}**: {cat['score']}/100")
            st.progress(cat["score"] / 100)
            if cat["feedback"]:
                st.caption(cat["feedback"])
        if result["strengths"]:
            st.subheader("Strengths")
            for s in result["strengths"]:
                st.markdown(f"- {s}")

    with tab_fix:
        if not result["improvements"]:
            st.info("No improvements returned.")
        icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
        for item in result["improvements"]:
            st.markdown(f"{icons[item['priority']]} **{item['priority']}**: {item['issue']}")
            if item["fix"]:
                st.markdown(f"> **Fix:** {item['fix']}")

    with tab_keys:
        if result["missing_keywords"]:
            st.write("Consider adding these (only where they are true for you):")
            st.markdown(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.success("No major missing keywords found.")

    with tab_rewrite:
        if not result["rewrite_examples"]:
            st.info("No rewrite examples returned.")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex['original']}")
            st.markdown(f"**After:** {ex['improved']}")
            st.divider()


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an estimated ATS score and concrete ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_config("GEMINI_API_KEY")
        if not api_key:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        model = st.text_input("Model", value=get_config("GEMINI_MODEL", DEFAULT_MODEL))
        st.caption(
            "Your resume is sent to Google's Gemini API for analysis and is not stored by this app."
        )

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste a job description to also get a job-match score and tailored keywords.",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar (or set GEMINI_API_KEY).")
            return
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            resume_text = extract_text(uploaded.name, uploaded.getvalue())
        except ImportError as exc:
            st.error(MISSING_DEPS_MSG.format(err=exc))
            return
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(resume_text) < MIN_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image or "
                "a design-heavy PDF, an ATS would struggle with it too. Export a text-based "
                "PDF or DOCX and try again."
            )
            return

        try:
            from google import genai

            with st.spinner("Analyzing your resume..."):
                client = genai.Client(api_key=api_key)
                result = analyze_resume(client, model.strip() or DEFAULT_MODEL, resume_text, job_description)
        except ImportError as exc:
            st.error(MISSING_DEPS_MSG.format(err=exc))
            return
        except ValueError as exc:  # unusable model output
            st.error(f"{exc} Please try again.")
            return
        except Exception as exc:  # network, quota, bad key, bad model name
            st.error(f"Gemini API error: {exc}")
            return

        render_results(result)
        with st.expander("Text extracted from your resume (what an ATS would see)"):
            st.text(resume_text[:MAX_CHARS])


if __name__ == "__main__":
    main()
